import fitz
from docx import Document
from pptx import Presentation
from PIL import Image
import pytesseract 
import os
from pathlib import Path

def extract_text(file_path: str, file_type: str)-> str:
    path =Path(file-path)
    if not path.exists():
        raise FileNotFoundError(f"File Not Found: {file_path}")
    if file_type == "pdf":
        return _parse_pdf (file_path)
    elif file_type == "docx":
        return _parse_docx (file_path)
    elif file_type == "pptx":
        return _parse_pptx (file_path)
    elif file_type == "image":
        return _parse_image (file_path)
    elif file_type =="txt":
        return _parse_txt(file_path)
    else: 
        raise ValueError(f"Unsupported file Type: {file_type}")

def _parse_pdf(path:str)-> str:
    doc = fitz.open(path)
    pages = []
    for page_num , page in enumerate(doc, start=1):
        text = page.get_text("text").script()
        if text:
            pages.append(f"[Page{page_num}]\n {text}")
    
    doc.close()
    return     "\n \n" . join (pages)


def _parse_docx(path: str) -> str:
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def _parse_pptx(path: str) -> str:
    prs = Presentation(path)
    slides = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"[Notes]: {notes}")
        if texts:
            slides.append(f"[Slide {slide_num}]\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _parse_image(path: str) -> str:
    img = Image.open(path)
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    text = pytesseract.image_to_string(img, lang="eng")
    return text.strip()

def _parse_txt(path:str)->str:
    with open(path, "r", encoding = "utf-8") as f:
        return f.read().strip()


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    type_map = {
        ".pdf":  "pdf",
        ".docx": "docx",
        ".doc":  "docx",
        ".pptx": "pptx",
        ".ppt":  "pptx",
        ".jpg":  "image",
        ".jpeg": "image",
        ".png":  "image",
        ".txt": "txt",
    }
    file_type = type_map.get(ext)
    if not file_type:
        raise ValueError(f"Unsupported file extension: {ext}")
    return file_type